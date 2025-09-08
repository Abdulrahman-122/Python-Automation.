# Gitting Rows and Culomns from the sheet:
#
# import openpyxl

# wb = openpyxl.load_workbook("duesRecords.xlsx")
# sheet = wb["Sheet1"]
# # print(tuple(sheet["A1":"H7"]))

# wb=openpyxl.load_workbook('duesRecords.xlsx')
# sh1=wb['Sheet1']
# Allcells=sh1['A1':'H7']
# for row_of_cell in Allcells:
# for cellobj in row_of_cell:
#     print(cellobj.coordinate,cellobj.value)
# print('---End of Row----')


#
# we use tuple to visualize the cells object  at the workbook object
#
# as you see the tuple of sheet contain on seven inner tubles
# each tuble represent the obj of each row from the left to the right
# so the objects into each row is the cell objects for each row.
# we use two for loops to print the cells values
# the outer goes for each each row
# the inner goes for each cell row


import openpyxl

wb = openpyxl.load_workbook("duesRecords.xlsx")
sh1 = wb["Sheet1"]
column_b = sh1["B"]
for collobj in column_b:
    print(collobj.value)
column_a = sh1["A"]
for collobj in column_a:
    print(collobj.value)
column_c = sh1["C"]
for collobj in column_c:
    print(collobj.value)

#each column
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
