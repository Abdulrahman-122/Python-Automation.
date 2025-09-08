# Creating and removing Sheets:
#
# import openpyxl

# wb = openpyxl.Workbook()

# sheet = wb.active
# print(sheet)
# sheet1 = wb.create_sheet(index=0)
# sheet2 = wb.create_sheet(index=2, title="Sheet2")
# sheet3 = wb.create_sheet(index=3, title="Sheet3")
# sheet4 = wb.create_sheet(index=4, title="sheet4")
# print(wb.sheetnames)
# wb.remove(wb["Sheet"])
# wb.remove(wb['Sheet1'])

# print(wb.sheetnames)

# wb.save('File1.xlsx')


# remove(worksheetobject)
# remove not take the sheetname just instead take worksheet object
#
# import openpyxl

# wb = openpyxl.load_workbook("File1.xlsx")
# print(wb.active)
# sh2 = wb["Sheet2"]
# for i in range(1, 100):
#     cells = sh2.cell(row=i, column=1)
#     cells.value = "Hello man!."
# print("Done.")
# print()
# wb.save("UpFile1.xlsx")
#
#
# Setting the font style of cell:

#
# to customize the font of text in excel: use font,style module from openpyxl

# import openpyxl
# from openpyxl.styles import Font

# wb = openpyxl.Workbook()
# sheet=wb.active
# italic24font=Font(size=25,italic=True)  # store inside italic24font (A font object with attributes (size and ittalic))
# sheet['A1'].font=italic24font
# sheet['A1']='Hello world!!!'
# wb.save('Font1.xlsx')
#
#write another word to the same excel file

import openpyxl
from openpyxl.styles import Font
wb=openpyxl.load_workbook('Font1.xlsx')

sheet=wb.active









#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
