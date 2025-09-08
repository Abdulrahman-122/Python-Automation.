# Multiplication Table Maker
#
#
# the program take N integer from the command and this will determine the size of table
# N*N  ex if you put:4 -> table will be 4*4 and so on
# -------------------------------------------


import openpyxl

wb = openpyxl.Workbook()
wb.create_sheet("Sheet1")
sheet = wb["Sheet1"]
Table_shape=int(input('Enter the integer to make (N*N) multiplication table:'))
for i in range(Table_shape+1):
    sheet.cell(row=i)
