# import openpyxl

# wb = openpyxl.load_workbook("example.xlsx")
# print("All sheets:", wb.sheetnames)
# first_sheet = wb["Sheet1"]
# print(first_sheet)
# Second_sheet = wb["Sheet2"]
# print(Second_sheet)
# Third_sheet = wb["Sheet3"]
# print(Third_sheet)
# print(first_sheet.title)
# print(Second_sheet.title)
# print(Third_sheet.title)
# another_sheet = wb.active
# print(another_sheet.title)


# EXcel sheet:
#
# openpyxl.load_workbook -> takes filename and return workbook data type
# wb.sheetnames -> return all sheets.
# this workbook data object is an  the EXcel file.
# wb.sheetnames -> return the sheets of the workbook
# wb.active -> return the top sheet when the workbook is opened in Excel.

# Getting Cells from the sheets:
# once you have a worksheet object
# to access the cell object by  name
#
#
#
import openpyxl

wb = openpyxl.load_workbook("example.xlsx")
# sheet = wb.sheetnames
sheet1 = wb["Sheet1"]
call_a1 = sheet1["A1"]
print(call_a1)
print(call_a1.value)

sheet1 = wb["Sheet1"]
cell_2 = sheet1["A2"]
print(cell_2)
print(cell_2.value)
sheet3 = wb["Sheet1"]
Cell_3 = sheet3["A3"]
print(Cell_3)
print(Cell_3.value)
sheet1 = wb["Sheet1"]
Cell_4 = sheet1["A4"]
print(Cell_4)
print(Cell_4.value)
sheet1 = wb["Sheet1"]
Cell_5 = sheet1["A5"]
print(Cell_5)
print(Cell_5.value)
sheet1 = wb["Sheet1"]
Cell_6 = sheet1["A6"]
print(Cell_6)
print(Cell_6.value)
sheet1 = wb["Sheet1"]
Cell_7 = sheet1["A7"]
print(Cell_7)
print(Cell_7.value)

# to access sheet 2
Sheet2 = wb["Sheet2"]
Cell1 = Sheet2["B1"]
Cell2 = Sheet2["B2"]
Cell3 = Sheet2["B3"]
Cell4 = Sheet2["B4"]
Cell5 = Sheet2["B5"]
Cell6 = Sheet2["B6"]
Cell7 = Sheet2["B7"]
print(Cell1.value)
print(Cell2.)
print(Cell3)
print(Cell4)
print(Cell5)
print(Cell6)
print(Cell7)
