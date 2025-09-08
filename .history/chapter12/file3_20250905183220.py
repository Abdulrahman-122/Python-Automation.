# Creating and removing Sheets:
#
import openpyxl

wb = openpyxl.Workbook()

sheet = wb.active
print(sheet)
sheet1=wb.create_sheet()
sheet2=wb.create_sheet(index=2,title='Sheet2')
sheet3=wb.create_sheet(index=3,title='Sheet')
# wb.save('File1.xlsx')


#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
