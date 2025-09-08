#Making Barchart contian on Months with Sales and profits:

# import openpyxl
# from openpyxl.chart import Reference,BarChart

# wb=openpyxl.Workbook()
# sheet=wb.active

# for i in range(1,9):
#     sheet['A'+str(i)]=i          # months
#     sheet['B'+str(i)]=i*2        # sales
#     sheet['C'+str(i)]=i          # profits
# sheet['A1']='Months'
# sheet['B1']='Sales'
# sheet['C1']='Profits'
# ref=Reference(sheet,min_col=2,max_col=3,min_row=1 ,max_row=8)       
# cats=Reference(sheet,min_col=1 ,min_row=1 ,max_row=8)
# Chart=BarChart()
# Chart.title='Sales Vs Profits'
# Chart.x_axis.name='Months'
# Chart.y_axis.name='Amount'

# Chart.add_data(ref,titles_from_data=True)
# Chart.set_categories(cats)
# sheet.add_chart(Chart,'E5')
# wb.save('Chart2.xlsx')

#Draw a chart with sales on y-axis and months on x-axis 
import openpyxl
from openpyxl.chart import Reference,BarChart

wb=openpyxl.Workbook(

)
sheet=wb.active
sheet['A2']