# Blank Row inserter:
# How it works:
#take two integers from the terminal : sys.argv N1 N2
# N1 -> the starting row
# N2 -> the number of blanck rows 
# ---
# you will need to spreadsheet to read it 
# then take two integers from user
# then write out the old spreadsheet into the new one
# then  by using for loop
# start at N1 row and put N2 rows blanck and write the rest of the old spread into the rest ones

import openpyxl,sys

try:
    N=int(sys.argv[1])
    M=int(sys.argv[2])
except:
    N=int(input('Enter the first number(starting Row):'))
    M=int(input('Enter the second number(Blanck Rows): '))

wb=openpyxl.load_workbook('produceSales.xlsx')
sheet=wb.active

for row in range(1,N+1):
    for col in range(1,sheet.max_column):
        sheet.cell(row=row,column=col).value=sheet.cell(row=row,column=col).value

for row in range(M+1,sheet.max_row):
    for col in range(1,sheet.max_column)






wb.save('Project4.xlsx')











# 
# 
# 
# 
# 
# 
# 
# 
