# Multiplication Table Maker
#
#
# the program take N integer from the command and this will determine the size of table
# N*N  ex if you put:4 -> table will be 4*4 and so on
# -------------------------------------------


import openpyxl,sys

wb=openpyxl.Workbook()
sheet=wb.active
try:
    N=int(sys.argv[1])
except:
    N=int(input('Enter an integer to make N*N Multiplication Table:'))




for i in range(1,N+1):
    sheet.cell(row=i+1,column=1).value=i #make column header in A column
    sheet.cell(row=1,column=i+1).value=i #make row headers

for row in range(2,N+2):
    for col in range(2,N+2):
        sheet.cell(row=row,column=col).value=sheet.cell(row=)