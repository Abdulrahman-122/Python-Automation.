# import openpyxl

# wb = openpyxl.load_workbook("example.xlsx")
# print("All sheets:", wb.sheetnames)
# first_sheet = wb["Sheet1"]
# print(first_sheet)
# Second_sheet = wb["Sheet2"]
# print(Second_sheet)
# Third_sheet = wb["Sheet3"]
# print(Third_sheet)
# print(first_sheet.title)
# print(Second_sheet.title)
# print(Third_sheet.title)
# another_sheet = wb.active
# print(another_sheet.title)


# EXcel sheet:
#
# openpyxl.load_workbook -> takes filename and return workbook data type
# wb.sheetnames -> return all sheets.
# this workbook data object is an  the EXcel file.
# wb.sheetnames -> return the sheets of the workbook
# wb.active -> return the top sheet when the workbook is opened in Excel.

# Getting Cells from the sheets:
# once you have a worksheet object
# to access the cell object by  name
#
#
#
import openpyxl

wb = openpyxl.load_workbook("example.xlsx")
# sheet = wb.sheetnames
sheet1 = wb['Sheet1']
call_a1=sheet1['A1']
print(call_a1)
print(call_a1.value)

sheet1=wb['Sheet1']
cell_2=sheet1['A2']
print(cell_2)
print(cell_2.value)
sheet3=wb['Sheet1']
Cell_3=sheet3['A3']
print(Cell_3)
print(Cell_3.value)
sheet1=wb['Sheet1']
Cell_4=she
