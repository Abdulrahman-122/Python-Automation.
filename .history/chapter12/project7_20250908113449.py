#write a spreadsheet into a textfile:
#
#
#
#
#
#
import openpyxl
wb=openpyxl.load_workbook('produceSales.xlsx')
sheet=wb.active
lines1=[]
# for row in range(1,sheet.max_row):
#     for col in range(1,sheet.max_column):
#         lines1.append(sheet.cell(row=row,column=col))

# file1=open('File7.txt','w')
# file1.writelines(lines1)
# print('Done')
columna=[]
for row in range(1,sheet.max_row):
    for col in range(1,sheet.max_column):
        sheet['A'+str(row)]