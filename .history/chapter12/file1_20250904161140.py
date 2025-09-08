# import openpyxl

# wb = openpyxl.load_workbook("example.xlsx")
# print("All sheets:", wb.sheetnames)
# first_sheet = wb["Sheet1"]
# print(first_sheet)
# Second_sheet = wb["Sheet2"]
# print(Second_sheet)
# Third_sheet = wb["Sheet3"]
# print(Third_sheet)
# print(first_sheet.title)
# print(Second_sheet.title)
# print(Third_sheet.title)
# another_sheet = wb.active    
# print(another_sheet.title)



# EXcel sheet:
#
# openpyxl.load_workbook -> takes filename and return workbook data type
# wb.sheetnames -> return all sheets.
# this workbook data object is an  the EXcel file.
# wb.sheetnames -> return the sheets of the workbook
# wb.active -> return the top sheet when the workbook is opened in Excel.

# Getting Cells from the sheets:
# once you have a worksheet object
# to access the cell object by  name 
# 
# 
# 
import openpyxl
wb=openpyxl.load_workbook('example.xlsx')
sheet=wb.sheetnames
sheet1=wb['sheet1']
