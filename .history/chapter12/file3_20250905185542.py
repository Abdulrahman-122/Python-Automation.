# Creating and removing Sheets:
#
# import openpyxl

# wb = openpyxl.Workbook()

# sheet = wb.active
# print(sheet)
# sheet1 = wb.create_sheet(index=0)
# sheet2 = wb.create_sheet(index=2, title="Sheet2")
# sheet3 = wb.create_sheet(index=3, title="Sheet3")
# sheet4 = wb.create_sheet(index=4, title="sheet4")
# print(wb.sheetnames)
# wb.remove(wb["Sheet"])
# wb.remove(wb['Sheet1'])

# print(wb.sheetnames)

# wb.save('File1.xlsx')


# remove(worksheetobject)
# remove not take the sheetname just instead take worksheet object
#
import openpyxl

wb = openpyxl.load_workbook("File1.xlsx")
print(wb.active)
sh2 = wb["Sheet2"]
sh2["A1"] = "Hello from python"
print(sh2["A1"].value)
for i in range(1,100):
    sh2.cell(row=i,column=1) =="Hello"
print("Done.")
print()
wb.save('.xlsx')

#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
