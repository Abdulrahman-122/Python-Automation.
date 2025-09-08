# import openpyxl

# wb = openpyxl.load_workbook("example.xlsx")
# print("All sheets:", wb.sheetnames)
# first_sheet = wb["Sheet1"]
# print(first_sheet)
# Second_sheet = wb["Sheet2"]
# print(Second_sheet)
# Third_sheet = wb["Sheet3"]
# print(Third_sheet)
# print(first_sheet.title)
# print(Second_sheet.title)
# print(Third_sheet.title)
# another_sheet = wb.active
# print(another_sheet.title)


# EXcel sheet:
#
# openpyxl.load_workbook -> takes filename and return workbook data type
# wb.sheetnames -> return all sheets.
# this workbook data object is an  the EXcel file.
# wb.sheetnames -> return the sheets of the workbook
# wb.active -> return the top sheet when the workbook is opened in Excel.

# Getting Cells from the sheets:
# once you have a worksheet object
# to access the cell object by  name
#
#
#
import openpyxl

wb = openpyxl.load_workbook("duesRecords.xlsx")
sheets = wb.sheetnames
print(sheets)
sh1 = wb["Sheet1"]
print(sh1)
c1 = sh1["A1"]
print(c1.value)
c2 = sh1["A2"]
print(c2.value)
c3 = sh1.cell(row=3, column=1)
print(c3.value)
sh2=wb['Sheet2']
print(sh2)
b1=sh1['B1']
print(b1.value)
b2=sh1['B2']
print(b2.value)
print(f'The {b1.value} of {c1.value} is = {b}')
