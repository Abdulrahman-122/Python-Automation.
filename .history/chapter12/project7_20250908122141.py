# write a spreadsheet into a textfile:
#
#
#
#
#
#
import openpyxl

wb = openpyxl.load_workbook("produceSales.xlsx",data_only=True)
sheet = wb.active
lines1 = []
# for row in range(1,sheet.max_row):
#     for col in range(1,sheet.max_column):
#         lines1.append(sheet.cell(row=row,column=col).value)
#         print('\n')

# file1=open('File7.txt','w')

# file1.writelines(lines1)

# print('Done')

columna = []
columnb = []
columnc = []
columnd=[]
for row in range(1, sheet.max_row):
    columna.append(sheet["A" + str(row)].value)
    columnb.append(sheet["B" + str(row)].value)
    columnc.append(sheet["C" + str(row)].value)
    columnd.append(sheet['D'+str(row)].value)
print(columna)
print(columnb)
print(columnc)
print(columnd)
with open("File7.txt", "w") as file1:
    for item1, item2, item3,item4 in zip(columna, columnb, columnc,columnd):
        # file1.write(item1 + " " + str(item2) + " " + item3  +" "+ str(item4) +'\n')
        file1.write(f"item1 + " " + str(item2) + " " + item3  +" "+ str(item4) +'\n'")



print("Done.")
 