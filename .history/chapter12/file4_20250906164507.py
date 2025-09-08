# Merging and Unmerging Cells:
# to merge a number of cells with each other
# import openpyxl
# wb=openpyxl.Workbook()
# sheet=wb.active
# sheet['A1']='Twelve merged cells with each other.'
# sheet.merge_cells('A1:D3')
# sheet['D5']='Two merged cells'
# sheet.merge_cells('D5:E5')
# wb.save('merge.xlsx')


import openpyxl

wb = openpyxl.load_workbook("merge.xlsx")
sheet = wb.active
sheet.unmerge_cells("A1:D3")
sheet.unmerge_cells('D5:E5')
wb.save("unmerge1.xlsx")
#
#Freeze panes:
#is to freeze a few of the top row and the leftmost column onscreen
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
