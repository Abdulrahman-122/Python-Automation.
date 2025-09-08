# Spreadsheet cell Inverter:
# 
# invert columns to rows and rows to columns
# 
# 
# to read use: sheetdata[x][y]inside for loop then 
# when write use; sheetdate[y][x] 
import openpyxl

wb=openpyxl.load_workbook('duesRecords.xlsx')
sheet=wb.active

New_wb=openpyxl.Workbook('Project5.xlsx')
New_sheet=New_wb.active

for x in range(1,sheet.max_row):
    for y in range(1,sheet.max_column):
        New_sheet.cell(row=y,)






# 
# 
# 
# 
# 
# 
# 
# 
# 
# 
# 
# 
# 
# 
# 
# 
# 
# 
# 
# 
# 
# 
