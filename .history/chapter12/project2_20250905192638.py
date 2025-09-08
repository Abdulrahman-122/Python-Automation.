# Updating a spreadsheet:
#
# in the project we will do:
# loop over the rows
# if the produce was galaric or celery or lemon-> update the price of them
# so;
# open the excel file
# if the row of produce has a cellvalue called galaric or celery or lemon
# update the price of them in column B
# save the old file to new one (in order to not lose the data inside it )
import openpyxl

wb = openpyxl.load_workbook("produceSales.xlsx")
print(wb.active)
print(wb.sheetnames)

sheet = wb["Sheet"]
print(sheet.max_row)
for i in range(1,23758):
    if sheet.cell(row=i,column=1).value=='Garlic':
        



#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
#
