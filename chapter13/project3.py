# Custom invitations as Word Documents:
# you have a text file : with a names of guests names
# one name in each line this file called: guests.txt
#
# make a program that : generate a custom invitaion
# contains on some sentenses
# it would be a pleasure to have the company of nameofguests
# at 11010 Memory Lane on the evening of April 1st at 7 o'clock
#
# what you will make
#
# add some styles to the blank doc like (the name of stylefont=Brush script std)
# and add italic to the whole sentenses except : name of guests
# while make nameofguests bold
# but underline under(at)
#
# then make one envitation for each name in the file (by making one page for each one and store all of theses pages inside one doc file)
# this make it ease: you will open one doc and print all invitaion at once
# instead of making one word for each invitation
#

import docx
from docx.text.run import WD_BREAK

doc = docx.Document()

with open("guests.txt", "r") as file:
    reader = file.readlines()
    # print(reader)
    content = [item.strip() for item in reader]
    # print(content)
    for name in content:
        doc.add_heading("Invitation Card", level=0)
        para = doc.add_paragraph("")
        run1 = para.add_run("It would be a pleasure to have the company of ")
        para.add_run("\n")
        run1.italic = True
        run2 = para.add_run(
            name,
        )
        run2.bold = True
        para.add_run("\n")
        run3 = para.add_run("at")
        run3.italic = True
        run3.underline = True
        run4 = para.add_run("11010 Memory Lane on the evening of")
        run4.italic = True
        para.add_run("\n")
        run40 = para.add_run("April 1st")
        run40.italic = True
        para.add_run("\n")
        run5 = para.add_run("at")
        run5.italic = True
        run5.underline = True
        run6 = para.add_run("7 o'clock")
        run6.italic = True
        if name != content[-1]:
            para.add_run("").add_break(WD_BREAK.PAGE)

    filename = input("Enter the name filename(.docx)>>")
    doc.save(filename)
print('Done.')